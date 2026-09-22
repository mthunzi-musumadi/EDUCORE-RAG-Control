# ==============================================================================
# TEST SUITE: DOCUMENT READER ABSTRACTION LAYER
# Unit tests for multi-format document extraction (DOCX, PDF, Excel)
# ==============================================================================
import sys
import os
import tempfile
import shutil
import unittest.mock

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'src', 'backend'))

from document_readers import (
    DocxReader,
    PdfReader,
    ExcelReader,
    read_document,
    get_format_prefix,
)


# ==============================================================================
# DOCX READER TESTS
# ==============================================================================

def test_docx_reader_returns_sections():
    """DocxReader correctly extracts heading-grouped sections from a .docx file."""
    import docx

    tmpdir = tempfile.mkdtemp()
    try:
        doc = docx.Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("This is the first paragraph of the introduction.")
        doc.add_paragraph("This is the second paragraph of the introduction.")

        file_path = os.path.join(tmpdir, "sample-doc.docx")
        doc.save(file_path)

        reader = DocxReader()
        clean_title, sections = reader.read(file_path)

        # Title derived from filename: "sample-doc" -> "Sample Doc"
        assert clean_title == "Sample Doc"

        # Should have at least one section with the heading we added
        assert len(sections) >= 1
        heading_names = [s["heading"] for s in sections]
        assert "Introduction" in heading_names

        intro_section = next(s for s in sections if s["heading"] == "Introduction")
        assert "first paragraph" in intro_section["content"]
        assert "second paragraph" in intro_section["content"]
    finally:
        shutil.rmtree(tmpdir)


# ==============================================================================
# PDF READER TESTS (pypdfium2 — 100% offline, zero-egress)
# ==============================================================================

def test_pdf_reader_extracts_from_fixture():
    """PdfReader parses pages from a real PDF file using pypdfium2 with zero egress."""
    fixture_path = os.path.join(os.path.dirname(__file__), "..", "fixtures", "sample_document.pdf")
    if not os.path.exists(fixture_path):
        pytest.skip("sample_document.pdf fixture not found")

    reader = PdfReader()
    # Test reading with page_range to keep test fast
    clean_title, sections = reader.read(fixture_path, page_range=(1, 3))

    assert clean_title == "Sample Document"
    assert len(sections) >= 1
    # Check that text was extracted from the fixture
    all_content = " ".join(s["content"] for s in sections)
    assert len(all_content) > 50


def test_pdf_reader_heading_detection():
    """PdfReader._extract_page_sections detects uppercase headings and sections."""
    page_text = (
        "SECTION 1 - GENERAL PROVISIONS\n"
        "This is the body of section one.\n"
        "It contains multiple lines.\n"
        "SECTION 2 - ELIGIBILITY\n"
        "Candidates must satisfy all criteria.\n"
    )
    sections = PdfReader._extract_page_sections(page_text, "Test Document", 1)

    assert len(sections) == 2
    headings = [s["heading"] for s in sections]
    assert any("SECTION 1" in h for h in headings)
    assert any("SECTION 2" in h for h in headings)


def test_pdf_reader_fallback_page_section():
    """PdfReader._extract_page_sections falls back to page title when no headings exist."""
    plain_text = "Just a standard paragraph without any uppercase or special section headings."
    sections = PdfReader._extract_page_sections(plain_text, "Report", 3)

    assert len(sections) == 1
    assert sections[0]["heading"] == "Report - Page 3"
    assert "standard paragraph" in sections[0]["content"]


# ==============================================================================
# EXCEL READER TESTS
# ==============================================================================

def test_excel_reader_multi_sheet():
    """ExcelReader produces one section per non-empty worksheet with pipe-delimited rows."""
    import openpyxl

    tmpdir = tempfile.mkdtemp()
    try:
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1.append(["Name", "Age"])
        ws1.append(["Alice", 30])
        ws1.append(["Bob", 25])

        ws2 = wb.create_sheet("Sheet2")
        ws2.append(["Country", "Code"])
        ws2.append(["Zambia", "ZM"])

        file_path = os.path.join(tmpdir, "data_file.xlsx")
        wb.save(file_path)
        wb.close()

        reader = ExcelReader()
        clean_title, sections = reader.read(file_path)

        assert clean_title == "Data File"
        assert len(sections) == 2

        # Sheet1 section
        s1 = next(s for s in sections if "Sheet1" in s["heading"])
        assert s1["heading"] == "Data File - Sheet1"
        assert "Name | Age" in s1["content"]
        assert "Alice | 30" in s1["content"]
        assert "Bob | 25" in s1["content"]

        # Sheet2 section
        s2 = next(s for s in sections if "Sheet2" in s["heading"])
        assert s2["heading"] == "Data File - Sheet2"
        assert "Zambia | ZM" in s2["content"]
    finally:
        shutil.rmtree(tmpdir)


def test_excel_reader_empty_sheet_skipped():
    """ExcelReader skips worksheets that contain no data rows."""
    import openpyxl

    tmpdir = tempfile.mkdtemp()
    try:
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Populated"
        ws1.append(["Key", "Value"])
        ws1.append(["alpha", "1"])

        # Create an empty sheet (no appended rows)
        wb.create_sheet("Empty")

        file_path = os.path.join(tmpdir, "mixed_sheets.xlsx")
        wb.save(file_path)
        wb.close()

        reader = ExcelReader()
        clean_title, sections = reader.read(file_path)

        assert clean_title == "Mixed Sheets"
        # Only the populated sheet should appear
        assert len(sections) == 1
        assert "Populated" in sections[0]["heading"]
    finally:
        shutil.rmtree(tmpdir)


# ==============================================================================
# READER REGISTRY & FACTORY TESTS
# ==============================================================================

def test_reader_registry_dispatch():
    """read_document dispatches .docx files to DocxReader and raises on unsupported extensions."""
    import docx

    tmpdir = tempfile.mkdtemp()
    try:
        # Create a minimal .docx and verify dispatch
        doc = docx.Document()
        doc.add_paragraph("Registry dispatch test content.")
        file_path = os.path.join(tmpdir, "dispatch_test.docx")
        doc.save(file_path)

        clean_title, sections = read_document(file_path)
        assert clean_title == "Dispatch Test"
        assert len(sections) >= 1

        # Unsupported extension should raise ValueError
        bad_path = os.path.join(tmpdir, "notes.txt")
        with open(bad_path, "w") as f:
            f.write("plain text")

        with pytest.raises(ValueError, match="Unsupported file format"):
            read_document(bad_path)
    finally:
        shutil.rmtree(tmpdir)


def test_unsupported_extension_raises():
    """read_document raises ValueError with a descriptive message for .txt files."""
    tmpdir = tempfile.mkdtemp()
    try:
        txt_path = os.path.join(tmpdir, "test.txt")
        with open(txt_path, "w") as f:
            f.write("hello")

        with pytest.raises(ValueError, match=r"Unsupported file format '\.txt'"):
            read_document(txt_path)
    finally:
        shutil.rmtree(tmpdir)


# ==============================================================================
# FORMAT PREFIX TESTS
# ==============================================================================

@pytest.mark.parametrize("filename, expected_prefix", [
    ("framework.docx", "EDU-FW"),
    ("report.pdf", "EDU-PDF"),
    ("data.xlsx", "EDU-XLS"),
    ("legacy.xls", "EDU-XLS"),
])
def test_get_format_prefix(filename, expected_prefix):
    """get_format_prefix returns the correct chunk-ID prefix for each supported extension."""
    assert get_format_prefix(filename) == expected_prefix
