# ==============================================================================
# TEST SUITE: MULTI-FORMAT DOCUMENT INGESTION INTEGRATION
# Tests PDF, Excel and DOCX discovery, parsing, classification and sync
# Compliance: ISO/IEC 42001:2023 | Zambian Data Protection Act No. 3 of 2021
# ==============================================================================
import os
import sys
import shutil
import tempfile
import docx
import openpyxl
import pytest

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
sys.path.insert(0, os.path.join(BASE_DIR, "src", "backend"))

from framework_sync_service import (
    FrameworkSyncService,
    SUPPORTED_EXTENSIONS,
)


def _create_docx(file_path: str, title: str, paragraphs: list):
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(file_path)


def _create_xlsx(file_path: str, sheet_data: dict):
    wb = openpyxl.Workbook()
    first = True
    for sheet_name, rows in sheet_data.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)
        for r in rows:
            ws.append(r)
    wb.save(file_path)
    wb.close()


def test_scan_discovers_pdf_and_xlsx():
    """scan_framework_files discovers .docx, .pdf, and .xlsx files across watch directories."""
    tmp_dir = tempfile.mkdtemp()
    try:
        # Create .docx
        _create_docx(os.path.join(tmp_dir, "policy.docx"), "Test Policy", ["Policy content"])
        # Create .xlsx
        _create_xlsx(os.path.join(tmp_dir, "data.xlsx"), {"Sheet1": [["Header1", "Header2"], ["Val1", "Val2"]]})
        # Create .pdf (copy from fixtures if available, else create dummy file)
        pdf_fixture = os.path.join(BASE_DIR, "tests", "fixtures", "sample_document.pdf")
        if os.path.exists(pdf_fixture):
            shutil.copy(pdf_fixture, os.path.join(tmp_dir, "handbook.pdf"))
        else:
            with open(os.path.join(tmp_dir, "handbook.pdf"), "wb") as f:
                f.write(b"%PDF-1.4 dummy pdf")

        svc = FrameworkSyncService(
            framework_dir=tmp_dir,
            state_file=os.path.join(tmp_dir, "state.json"),
            data_file=os.path.join(tmp_dir, "data.json"),
        )
        scanned = svc.scan_framework_files()

        scanned_filenames = [info["filename"] for info in scanned.values()]
        assert "policy.docx" in scanned_filenames
        assert "data.xlsx" in scanned_filenames
        assert "handbook.pdf" in scanned_filenames
        assert len(scanned) == 3
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_scan_ignores_unsupported_formats():
    """scan_framework_files ignores unsupported formats like .txt and .csv."""
    tmp_dir = tempfile.mkdtemp()
    try:
        _create_docx(os.path.join(tmp_dir, "valid.docx"), "Valid Doc", ["Text"])
        with open(os.path.join(tmp_dir, "notes.txt"), "w") as f:
            f.write("text content")
        with open(os.path.join(tmp_dir, "data.csv"), "w") as f:
            f.write("col1,col2\nval1,val2")

        svc = FrameworkSyncService(
            framework_dir=tmp_dir,
            state_file=os.path.join(tmp_dir, "state.json"),
            data_file=os.path.join(tmp_dir, "data.json"),
        )
        scanned = svc.scan_framework_files()

        scanned_filenames = [info["filename"] for info in scanned.values()]
        assert "valid.docx" in scanned_filenames
        assert "notes.txt" not in scanned_filenames
        assert "data.csv" not in scanned_filenames
        assert len(scanned) == 1
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_scan_ignores_temp_files():
    """scan_framework_files ignores temporary lock files starting with ~$ or ."""
    tmp_dir = tempfile.mkdtemp()
    try:
        _create_docx(os.path.join(tmp_dir, "real.docx"), "Real Doc", ["Text"])
        with open(os.path.join(tmp_dir, "~$real.docx"), "w") as f:
            f.write("lock")
        with open(os.path.join(tmp_dir, ".hidden.xlsx"), "w") as f:
            f.write("hidden")

        svc = FrameworkSyncService(
            framework_dir=tmp_dir,
            state_file=os.path.join(tmp_dir, "state.json"),
            data_file=os.path.join(tmp_dir, "data.json"),
        )
        scanned = svc.scan_framework_files()

        scanned_filenames = [info["filename"] for info in scanned.values()]
        assert scanned_filenames == ["real.docx"]
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_parse_xlsx_to_records():
    """parse_file_to_records generates structured records with EDU-XLS prefix for Excel files."""
    tmp_dir = tempfile.mkdtemp()
    try:
        xlsx_path = os.path.join(tmp_dir, "student-records.xlsx")
        _create_xlsx(xlsx_path, {
            "Term1": [
                ["Student ID", "Name", "Grade"],
                ["ST-001", "Alice Tembo", "A"],
                ["ST-002", "Bob Mwape", "B"]
            ]
        })

        svc = FrameworkSyncService(
            framework_dir=tmp_dir,
            state_file=os.path.join(tmp_dir, "state.json"),
            data_file=os.path.join(tmp_dir, "data.json"),
        )
        scanned = svc.scan_framework_files()
        file_info = scanned["student-records.xlsx"]

        records = svc.parse_file_to_records(file_info)
        assert len(records) >= 1
        record = records[0]

        # Chunk ID format verification
        assert record["id"].startswith("EDU-XLS-")
        assert "Term1" in record["title"]
        assert "Alice Tembo" in record["content"]
        assert "Bob Mwape" in record["content"]
        assert record["source_file"] == "student-records.xlsx"
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_parse_docx_to_records_preserves_prefix():
    """parse_file_to_records maintains backward compatibility with EDU-FW- prefix for .docx files."""
    tmp_dir = tempfile.mkdtemp()
    try:
        docx_path = os.path.join(tmp_dir, "governance-policy.docx")
        _create_docx(docx_path, "Governance Policy", ["Section 1 content for testing."])

        svc = FrameworkSyncService(
            framework_dir=tmp_dir,
            state_file=os.path.join(tmp_dir, "state.json"),
            data_file=os.path.join(tmp_dir, "data.json"),
        )
        scanned = svc.scan_framework_files()
        file_info = scanned["governance-policy.docx"]

        records = svc.parse_file_to_records(file_info)
        assert len(records) >= 1
        assert records[0]["id"].startswith("EDU-FW-")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_classification_excel_finance():
    """Excel files with financial keywords auto-escalate to admin clearance and finance category."""
    tmp_dir = tempfile.mkdtemp()
    try:
        xlsx_path = os.path.join(tmp_dir, "school-budget-2026.xlsx")
        _create_xlsx(xlsx_path, {
            "Budget": [
                ["Department", "Payroll Allocation (ZMW)", "Status"],
                ["Science Lab", "150000", "Approved"],
                ["Administration", "250000", "Pending"]
            ]
        })

        svc = FrameworkSyncService(
            framework_dir=tmp_dir,
            state_file=os.path.join(tmp_dir, "state.json"),
            data_file=os.path.join(tmp_dir, "data.json"),
        )
        scanned = svc.scan_framework_files()
        file_info = scanned["school-budget-2026.xlsx"]

        records = svc.parse_file_to_records(file_info)
        assert len(records) >= 1
        record = records[0]

        assert record["clearance"] == "admin"
        assert record["category"] == "finance"
        assert "ADMIN / FINANCE" in record["classification"]
        assert "admin" in record["allowed_roles"]
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_incremental_sync_mixed_formats():
    """generate_incremental_update accurately handles mixed-format batches."""
    tmp_dir = tempfile.mkdtemp()
    try:
        _create_docx(os.path.join(tmp_dir, "handbook.docx"), "Staff Handbook", ["Rules and regulations."])
        _create_xlsx(os.path.join(tmp_dir, "inventory.xlsx"), {"Items": [["Item", "Quantity"], ["Desk", "50"]]})

        state_file = os.path.join(tmp_dir, "state.json")
        data_file = os.path.join(tmp_dir, "data.json")

        svc = FrameworkSyncService(
            framework_dir=tmp_dir,
            state_file=state_file,
            data_file=data_file,
        )

        delta = svc.generate_incremental_update(force=True)

        assert delta["changed"] is True
        assert len(delta["records_to_upsert"]) >= 2
        assert any("handbook.docx" in f for f in delta["updated_files"])
        assert any("inventory.xlsx" in f for f in delta["updated_files"])
        assert os.path.exists(state_file)
        assert os.path.exists(data_file)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
