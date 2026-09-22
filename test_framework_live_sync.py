# ==============================================================================
# TEST SUITE: REAL-TIME FRAMEWORK DOCUMENT LIVE SYNC & EMBEDDINGS
# Tests dynamic change detection, incremental upsert, RAG retrieval, and cleanup
# Compliance: ISO/IEC 42001:2023 | Zambian Data Protection Act No. 3 of 2021
# ==============================================================================
import os
import sys
import time
import json
import subprocess
import urllib.request
import urllib.error
import docx

TEST_PORT = 8002
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PYTHON_EXE = os.path.join(BASE_DIR, "framework_control", "Scripts", "python.exe")
SERVER_SCRIPT = os.path.join(BASE_DIR, "educore_enterprise_backend.py")
TEST_DOCX_PATH = os.path.join(
    BASE_DIR, "EDUCORE_AI_FRAMEWORK", "01_SPOKES_POLICIES", "educore-test-live-sync-policy.docx"
)

SERVER_PROC = None

def start_server() -> bool:
    global SERVER_PROC
    print(f"Starting test backend server on port {TEST_PORT}...")
    SERVER_PROC = subprocess.Popen([PYTHON_EXE, SERVER_SCRIPT, str(TEST_PORT)])
    for _ in range(30):
        time.sleep(0.5)
        try:
            with urllib.request.urlopen(f"http://127.0.0.1:{TEST_PORT}/health", timeout=1) as resp:
                if resp.status == 200:
                    print("Server online.")
                    return True
        except Exception:
            pass
    return False

def stop_server():
    global SERVER_PROC
    if SERVER_PROC:
        print("Stopping test backend server...")
        SERVER_PROC.terminate()
        try:
            SERVER_PROC.wait(timeout=3)
        except Exception:
            SERVER_PROC.kill()

def create_docx(path: str, title: str, paragraphs: list):
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(path)

def test_live_sync_lifecycle():
    base_url = f"http://127.0.0.1:{TEST_PORT}"

    try:
        # 1. Test /api/framework/status
        print("\n--- 1. Testing GET /api/framework/status ---")
        with urllib.request.urlopen(f"{base_url}/api/framework/status") as resp:
            assert resp.status == 200
            status_data = json.loads(resp.read().decode("utf-8"))
            print("Status response:", json.dumps(status_data, indent=2))
            assert status_data["status"] == "online"
            initial_vector_count = status_data["total_vector_count"]
            print(f"Initial vector count: {initial_vector_count}")

        # 2. Test initial sync (should be up_to_date)
        print("\n--- 2. Testing POST /api/framework/sync (no-op check) ---")
        req = urllib.request.Request(
            f"{base_url}/api/framework/sync",
            data=json.dumps({}).encode("utf-8"),
            headers={"Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req) as resp:
            sync_data = json.loads(resp.read().decode("utf-8"))
            print("Sync response:", sync_data)
            assert resp.status == 200

        # 3. Create a brand new framework .docx policy file
        print("\n--- 3. Creating new test framework document ---")
        unique_secret_code = "MODULE-Q99-QUANTUM"
        create_docx(
            TEST_DOCX_PATH,
            "Educore Advanced Robotics and Quantum Simulator Policy",
            [
                f"SECTION 1 - QUALIFICATION PREREQUISITE: Students must successfully complete {unique_secret_code} certification before operating high-voltage quantum simulation rigs at Sentinel Kabitaka campus.",
                "SECTION 2 - SAFETY PROTOCOL: Any unauthorized tampering with lab hardware results in immediate revocation of campus laboratory clearance."
            ]
        )
        print(f"Created: {TEST_DOCX_PATH}")

        # 4. Trigger sync to detect and embed the new file
        print("\n--- 4. Triggering POST /api/framework/sync for new file ---")
        req = urllib.request.Request(
            f"{base_url}/api/framework/sync",
            data=json.dumps({}).encode("utf-8"),
            headers={"Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req) as resp:
            sync_data = json.loads(resp.read().decode("utf-8"))
            print("Sync response after new file:", sync_data)
            assert sync_data["changed"] is True
            assert sync_data["records_upserted"] > 0
            assert any("test-live-sync" in f for f in sync_data["updated_files"])

        # 5. Query RAG to verify retrieval of newly embedded policy
        print("\n--- 5. Querying RAG endpoint to verify retrieval ---")
        chat_req = urllib.request.Request(
            f"{base_url}/v1/chat/completions",
            data=json.dumps({
                "model": "educore-enterprise-all",
                "messages": [
                    {"role": "user", "content": "What module must students complete before operating high-voltage quantum simulation rigs at Sentinel Kabitaka?"}
                ],
                "stream": False
            }).encode("utf-8"),
            headers={"Content-Type": "application/json", "X-Educore-User": "staff"}
        )
        with urllib.request.urlopen(chat_req, timeout=90) as resp:
            chat_data = json.loads(resp.read().decode("utf-8"))
            content = chat_data["choices"][0]["message"]["content"]
            print("RAG Answer:\n", content)
            assert unique_secret_code in content or "Q99" in content or "Quantum" in content
            print("  [PASS] Newly embedded content was retrieved and answered accurately!")

        # 6. Update document with modified content
        print("\n--- 6. Updating document content (Simulating policy revision) ---")
        updated_code = "MODULE-Q101-SUPERCONDUCTOR"
        create_docx(
            TEST_DOCX_PATH,
            "Educore Advanced Robotics and Quantum Simulator Policy",
            [
                f"SECTION 1 - REVISED PREREQUISITE: Under revised 2026 guidelines, students must now complete {updated_code} certification before operating high-voltage quantum simulation rigs at Sentinel Kabitaka campus.",
                "SECTION 2 - SAFETY PROTOCOL: Any unauthorized tampering with lab hardware results in immediate revocation of campus laboratory clearance."
            ]
        )

        # 7. Trigger sync again
        print("\n--- 7. Triggering sync for revised document ---")
        with urllib.request.urlopen(req) as resp:
            sync_data = json.loads(resp.read().decode("utf-8"))
            print("Sync response after revision:", sync_data)
            assert sync_data["changed"] is True
            assert sync_data["records_upserted"] > 0

        # 8. Query again and assert updated content is reflected
        print("\n--- 8. Querying RAG endpoint for updated policy ---")
        with urllib.request.urlopen(chat_req, timeout=90) as resp:
            chat_data = json.loads(resp.read().decode("utf-8"))
            content = chat_data["choices"][0]["message"]["content"]
            print("Updated RAG Answer:\n", content)
            assert updated_code in content or "Q101" in content
            print("  [PASS] Updated document embedding was immediately reflected in RAG answer!")

        # 9. Cleanup: Delete test file and verify eviction
        print("\n--- 9. Deleting test document and verifying vector eviction ---")
        if os.path.exists(TEST_DOCX_PATH):
            os.remove(TEST_DOCX_PATH)

        with urllib.request.urlopen(req) as resp:
            sync_data = json.loads(resp.read().decode("utf-8"))
            print("Sync response after deletion:", sync_data)
            assert sync_data["changed"] is True
            assert sync_data["ids_deleted"] > 0
            print("  [PASS] Outdated document vectors cleanly evicted from Chroma store!")

        print("\n=== ALL REAL-TIME DOCUMENT LIVE SYNC TESTS PASSED! ===")

    finally:
        if os.path.exists(TEST_DOCX_PATH):
            try:
                os.remove(TEST_DOCX_PATH)
            except Exception:
                pass

if __name__ == "__main__":
    if start_server():
        try:
            test_live_sync_lifecycle()
        finally:
            stop_server()
    else:
        print("Failed to start server.")
        sys.exit(1)
